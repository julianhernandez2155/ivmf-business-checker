"""
Business Operational Status Checker — GUI
Double-click this file to launch. No terminal needed.

First-time setup: see SETUP.md
"""

import importlib.util
import os
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime

import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from concurrent.futures import ThreadPoolExecutor, as_completed

# ── Dependency check — runs before any third-party imports ────────────────────
# Uses importlib.util.find_spec so tkinter can still show the error dialog even
# if the missing package would have caused an ImportError on the next line.

def _check_deps():
    checks = {
        "dotenv":    "python-dotenv",
        "requests":  "requests",
        "pydantic":  "pydantic",
        "openpyxl":  "openpyxl",
        "bs4":       "beautifulsoup4",
    }
    return [pkg for mod, pkg in checks.items() if importlib.util.find_spec(mod) is None]

_missing = _check_deps()
if _missing:
    _root = tk.Tk()
    _root.withdraw()
    messagebox.showerror(
        "Missing Dependencies",
        f"Required packages not installed:\n\n  {', '.join(_missing)}\n\n"
        "Open a terminal in this folder and run:\n\n  pip install -r requirements.txt"
    )
    sys.exit(1)

# Third-party imports — safe after dep check
from dotenv import load_dotenv
import openpyxl

# Load .env from the same folder as this script
load_dotenv(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"))

from tools.check_business import check_business
from tools.checkpoint import (
    load_checkpoint, save_checkpoint, rewrite_checkpoint,
    get_run_summary, ERROR_KEYWORDS,
)
from tools.build_output import build_output_excel
from tools.columns import detect_columns, COST_PER_BUSINESS_ESTIMATE

# ── Constants ─────────────────────────────────────────────────────────────────

LOG_MAX_LINES = 500
DEFAULT_WORKERS = 1

WORKER_OPTIONS = [
    ("1 worker  (Tier 0 — safe for new accounts)", 1),
    ("3 workers (Tier 1 — faster)",                3),
    ("8 workers (Tier 2 — fastest)",               8),
]

# ── Helpers ───────────────────────────────────────────────────────────────────

def get_cell_value(ws, row_idx, col_idx):
    if col_idx is None:
        return ""
    return str(ws.cell(row_idx, col_idx + 1).value or "").strip()

def open_file(path):
    """Open a file with the default application (cross-platform)."""
    try:
        if sys.platform == "win32":
            os.startfile(path)
        elif sys.platform == "darwin":
            subprocess.call(["open", path])
        else:
            subprocess.call(["xdg-open", path])
    except Exception:
        pass


# ── Summary Dialog ────────────────────────────────────────────────────────────

class SummaryDialog(tk.Toplevel):
    def __init__(self, parent, summary, on_export, on_retry):
        super().__init__(parent)
        self.title("Run Complete — Summary")
        self.geometry("420x380")
        self.configure(bg="#1e2a3a")
        self.resizable(False, False)
        self.grab_set()

        total  = summary["Total"]
        active = summary["Active"]
        closed = summary["Likely Closed"]
        uncert = summary["Uncertain"]
        no_web = summary["No Web Presence"]
        errors = summary["Errors"]

        tk.Label(self, text="Run Complete!", font=("Arial", 14, "bold"),
                 fg="#4ddb8a", bg="#1e2a3a").pack(pady=(18, 4))
        tk.Label(self, text="Here's a breakdown of results:",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1e2a3a").pack()
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=12)

        stats = tk.Frame(self, bg="#0d1b2a", padx=20, pady=14)
        stats.pack(fill="x", padx=20)

        def pct(n):
            return f"({int(n / total * 100)}%)" if total > 0 else "(0%)"

        rows = [
            ("✓  Active",           str(active),  "#4ddb8a", pct(active)),
            ("✗  Likely Closed",    str(closed),  "#ff6b6b", pct(closed)),
            ("?  Uncertain",        str(uncert),  "#f0c040", pct(uncert)),
            ("○  No Web Presence",  str(no_web),  "#a0b8cc", pct(no_web)),
            ("⚠  Errors / Retries", str(errors),  "#ff9f4a", pct(errors)),
            ("─" * 18,              "",            "#2e4057", ""),
            ("Total Checked",       str(total),   "#c8d8e8", ""),
        ]
        for i, (label, val, color, pct_str) in enumerate(rows):
            tk.Label(stats, text=label, font=("Arial", 10), fg=color,
                     bg="#0d1b2a", anchor="w", width=22).grid(row=i, column=0, sticky="w", pady=2)
            tk.Label(stats, text=val, font=("Arial", 10, "bold"), fg=color,
                     bg="#0d1b2a", anchor="e", width=6).grid(row=i, column=1, sticky="e", pady=2)
            tk.Label(stats, text=pct_str, font=("Arial", 9), fg="#5a7a95",
                     bg="#0d1b2a", anchor="w", width=8).grid(row=i, column=2, sticky="w", padx=(6, 0), pady=2)

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=12)

        btn = tk.Frame(self, bg="#1e2a3a")
        btn.pack(pady=(0, 16))

        tk.Button(btn, text="💾  Export Results", font=("Arial", 10, "bold"),
                  bg="#1a4a6b", fg="white", relief="flat", padx=14, pady=6,
                  cursor="hand2", command=lambda: [self.destroy(), on_export()]
                  ).pack(side="left", padx=6)

        retry = tk.Button(btn, text=f"🔁  Retry {errors} Failed", font=("Arial", 10),
                          bg="#6b3a1a", fg="white", relief="flat", padx=14, pady=6,
                          cursor="hand2", command=lambda: [self.destroy(), on_retry()])
        retry.pack(side="left", padx=6)
        if errors == 0:
            retry.config(state="disabled", bg="#2e4057", fg="#5a7a95")

        tk.Button(btn, text="Close", font=("Arial", 10),
                  bg="#2e4057", fg="#a0b8cc", relief="flat", padx=14, pady=6,
                  cursor="hand2", command=self.destroy).pack(side="left", padx=6)


# ── Column Confirmation Dialog ────────────────────────────────────────────────

class ColumnConfirmDialog(tk.Toplevel):
    def __init__(self, parent, headers, detected, total_rows=0, already_done=0, workers=1):
        super().__init__(parent)
        self.title("Confirm Column Mapping")
        self.geometry("540x480")
        self.configure(bg="#1e2a3a")
        self.resizable(False, False)
        self.grab_set()

        self.headers            = headers
        self.result             = None
        self.chosen_workers     = workers
        self.vars               = {}
        self.selected_workers   = tk.IntVar(value=workers)

        tk.Label(self, text="Confirm Column Mapping",
                 font=("Arial", 13, "bold"), fg="#4da6ff", bg="#1e2a3a").pack(
                     pady=(16, 4), padx=20, anchor="w")
        tk.Label(self, text="The program detected the following columns. Correct any mistakes before starting.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1e2a3a", wraplength=500, justify="left"
                 ).pack(padx=20, anchor="w")
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        grid = tk.Frame(self, bg="#1e2a3a")
        grid.pack(padx=24, fill="x")

        header_options     = [f"{i+1}: {h}" for i, h in enumerate(headers) if h is not None]
        header_options_map = {f"{i+1}: {h}": i for i, h in enumerate(headers) if h is not None}
        self.header_options_map = header_options_map

        fields = [
            ("Business Name column:", "name"),
            ("Website column:",       "website"),
            ("City column:",          "city"),
            ("State column:",         "state"),
        ]
        for row_i, (label, field) in enumerate(fields):
            tk.Label(grid, text=label, font=("Arial", 10), fg="#a0b8cc",
                     bg="#1e2a3a", anchor="w", width=24).grid(row=row_i, column=0, pady=6, sticky="w")
            var     = tk.StringVar()
            det_idx = detected.get(field)
            if det_idx is not None and det_idx < len(headers) and headers[det_idx] is not None:
                var.set(f"{det_idx+1}: {headers[det_idx]}")
            elif header_options:
                var.set(header_options[0])
            cb = ttk.Combobox(grid, textvariable=var, values=header_options,
                              state="readonly", width=32, font=("Arial", 10))
            cb.grid(row=row_i, column=1, pady=6, padx=(8, 0), sticky="w")
            self.vars[field] = var

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        # ── Run estimate + worker selector ────────────────────────────────
        remaining    = max(0, total_rows - already_done)
        cost_per_biz = COST_PER_BUSINESS_ESTIMATE

        est_frame = tk.Frame(self, bg="#0d1b2a", padx=14, pady=10)
        est_frame.pack(fill="x", padx=20, pady=(0, 6))

        tk.Label(est_frame, text="Run Estimate", font=("Arial", 9, "bold"),
                 fg="#4da6ff", bg="#0d1b2a").grid(row=0, column=0, columnspan=6, sticky="w", pady=(0, 6))

        def stat(parent, label, value, color, row, col):
            tk.Label(parent, text=label, font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a").grid(
                row=row, column=col*2, sticky="w", padx=(0, 4))
            tk.Label(parent, text=value, font=("Arial", 10, "bold"), fg=color, bg="#0d1b2a").grid(
                row=row, column=col*2+1, sticky="w", padx=(0, 20))

        stat(est_frame, "Total rows:",   str(total_rows),   "#c8d8e8", 1, 0)
        stat(est_frame, "Already done:", str(already_done), "#4ddb8a", 1, 1)
        stat(est_frame, "Remaining:",    str(remaining),    "#f0c040", 1, 2)

        tk.Label(est_frame, text="Est. cost:",  font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a"
                 ).grid(row=2, column=0, sticky="w", padx=(0, 4))
        tk.Label(est_frame, text="Est. time:",  font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a"
                 ).grid(row=2, column=2, sticky="w", padx=(0, 4))
        tk.Label(est_frame, text="Workers:",    font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a"
                 ).grid(row=2, column=4, sticky="w", padx=(0, 4))

        cost_label    = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#ff9f4a", bg="#0d1b2a")
        time_label    = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#4da6ff", bg="#0d1b2a")
        workers_label = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#c8d8e8", bg="#0d1b2a")
        cost_label.grid(   row=2, column=1, sticky="w", padx=(0, 20))
        time_label.grid(   row=2, column=3, sticky="w", padx=(0, 20))
        workers_label.grid(row=2, column=5, sticky="w")

        def update_estimates(*_):
            w         = self.selected_workers.get()
            # Assume ~4s average API latency per call
            est_secs  = (remaining / w) * 4 if w > 0 else remaining * 4
            est_hours = est_secs / 3600
            est_time  = f"{est_hours:.1f} hrs" if est_hours >= 1 else f"{int(est_secs / 60)} min"
            cost_label.config(   text=f"~${remaining * cost_per_biz:.2f}")
            time_label.config(   text=est_time)
            workers_label.config(text=str(w))

        # Worker selector
        worker_frame = tk.Frame(est_frame, bg="#0d1b2a")
        worker_frame.grid(row=3, column=0, columnspan=6, sticky="w", pady=(10, 0))
        tk.Label(worker_frame, text="Workers:", font=("Arial", 9), fg="#a0b8cc",
                 bg="#0d1b2a").pack(side="left", padx=(0, 8))
        for label, val in WORKER_OPTIONS:
            tk.Radiobutton(
                worker_frame, text=label, variable=self.selected_workers, value=val,
                font=("Arial", 9), fg="#a0b8cc", bg="#0d1b2a",
                selectcolor="#1e2a3a", activebackground="#0d1b2a",
                command=update_estimates,
            ).pack(side="left", padx=(0, 12))

        update_estimates()
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        btn_frame = tk.Frame(self, bg="#1e2a3a")
        btn_frame.pack(pady=(0, 16))
        tk.Button(btn_frame, text="✓  Confirm & Start", font=("Arial", 10, "bold"),
                  bg="#1a6b3a", fg="white", relief="flat", padx=16, pady=6,
                  cursor="hand2", command=self._confirm).pack(side="left", padx=8)
        tk.Button(btn_frame, text="Cancel", font=("Arial", 10),
                  bg="#2e4057", fg="#a0b8cc", relief="flat", padx=16, pady=6,
                  cursor="hand2", command=self.destroy).pack(side="left", padx=8)

    def _confirm(self):
        mapping = {}
        for field, var in self.vars.items():
            val = var.get()
            if val in self.header_options_map:
                mapping[field] = self.header_options_map[val]
            else:
                messagebox.showerror("Error", f"Please select a valid column for: {field}")
                return
        self.result         = mapping
        self.chosen_workers = self.selected_workers.get()
        self.destroy()


# ── Main App ──────────────────────────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Business Operational Status Checker")
        self.geometry("940x700")
        self.resizable(True, True)
        self.configure(bg="#1e2a3a")

        # State
        self.api_key      = tk.StringVar(value=os.getenv("PERPLEXITY_API_KEY", ""))
        self.run_dir      = None
        self.col_map      = None
        self.num_workers  = DEFAULT_WORKERS
        self.stop_flag    = False
        self.pause_flag   = False
        self.running      = False
        self._log_lines   = 0
        self._throughput  = []  # timestamps of completed rows for ETA
        self._total_rows  = 0
        self._done_at_start = 0
        self._session_cost    = 0.0
        self._session_checked = 0

        self._build_ui()
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    # ── UI Build ──────────────────────────────────────────────────────────────

    def _build_ui(self):
        # Header
        hdr = tk.Frame(self, bg="#1e2a3a", pady=14)
        hdr.pack(fill="x", padx=20)
        tk.Label(hdr, text="Business Operational Status Checker",
                 font=("Arial", 16, "bold"), fg="#4da6ff", bg="#1e2a3a").pack(anchor="w")
        tk.Label(hdr, text="IVMF  |  Powered by Perplexity Sonar AI",
                 font=("Arial", 10), fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w")
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20)

        # API Key row
        api_frame = tk.Frame(self, bg="#1e2a3a", pady=12)
        api_frame.pack(fill="x", padx=20)
        tk.Label(api_frame, text="Perplexity API Key:", font=("Arial", 10),
                 fg="#a0b8cc", bg="#1e2a3a", width=20, anchor="w").grid(row=0, column=0, sticky="w")
        self.api_entry = tk.Entry(api_frame, textvariable=self.api_key, show="*",
                                  font=("Arial", 10), bg="#2a3f55", fg="white",
                                  insertbackground="white", relief="flat", width=48)
        self.api_entry.grid(row=0, column=1, sticky="ew", padx=(0, 8), ipady=5)
        tk.Button(api_frame, text="Show", font=("Arial", 9), bg="#2e4057", fg="#a0b8cc",
                  relief="flat", cursor="hand2",
                  command=lambda: self.api_entry.config(
                      show="" if self.api_entry.cget("show") == "*" else "*")
                  ).grid(row=0, column=2, padx=(0, 4))
        tk.Button(api_frame, text="Test Key", font=("Arial", 9), bg="#1a4a6b", fg="white",
                  relief="flat", cursor="hand2", command=self._test_api_key
                  ).grid(row=0, column=3, padx=(0, 4))

        # Key status line
        env_loaded = bool(os.getenv("PERPLEXITY_API_KEY"))
        self.key_status = tk.Label(
            api_frame,
            text="✓ Key loaded from .env" if env_loaded else "No key found — enter manually or add to .env",
            font=("Arial", 8),
            fg="#4ddb8a" if env_loaded else "#f0c040",
            bg="#1e2a3a"
        )
        self.key_status.grid(row=1, column=1, sticky="w", pady=(2, 0))
        api_frame.columnconfigure(1, weight=1)
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20)

        # Run mode cards
        run_frame = tk.Frame(self, bg="#1e2a3a", pady=14)
        run_frame.pack(fill="x", padx=20)
        tk.Label(run_frame, text="Select a run mode to begin:",
                 font=("Arial", 10), fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w", pady=(0, 8))
        modes = tk.Frame(run_frame, bg="#1e2a3a")
        modes.pack(fill="x")

        new_card = tk.Frame(modes, bg="#1a3a52", padx=16, pady=12)
        new_card.pack(side="left", fill="both", expand=True, padx=(0, 8))
        tk.Label(new_card, text="▶  New Run", font=("Arial", 11, "bold"),
                 fg="#4da6ff", bg="#1a3a52").pack(anchor="w")
        tk.Label(new_card, text="Pick any Excel file to start\na fresh check from scratch.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1a3a52", justify="left").pack(anchor="w", pady=(4, 8))
        tk.Button(new_card, text="Browse for Excel file", font=("Arial", 10),
                  bg="#1a6b3a", fg="white", relief="flat", padx=12, pady=6,
                  cursor="hand2", command=self._new_run).pack(anchor="w")

        res_card = tk.Frame(modes, bg="#1a3a52", padx=16, pady=12)
        res_card.pack(side="left", fill="both", expand=True)
        tk.Label(res_card, text="⏩  Resume Run", font=("Arial", 11, "bold"),
                 fg="#f0c040", bg="#1a3a52").pack(anchor="w")
        tk.Label(res_card, text="Pick a previous run folder inside\n'Runs/' to continue where you left off.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1a3a52", justify="left").pack(anchor="w", pady=(4, 8))
        tk.Button(res_card, text="Browse for run folder", font=("Arial", 10),
                  bg="#6b5a1a", fg="white", relief="flat", padx=12, pady=6,
                  cursor="hand2", command=self._resume_run).pack(anchor="w")

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=(10, 0))

        # Run info + progress
        self.run_info = tk.Label(self, text="No run loaded.",
                                 font=("Arial", 9, "italic"), fg="#4a6a85", bg="#1e2a3a")
        self.run_info.pack(anchor="w", padx=22, pady=(6, 0))

        prog_frame = tk.Frame(self, bg="#1e2a3a")
        prog_frame.pack(fill="x", padx=20, pady=(4, 0))
        self.progress_label = tk.Label(prog_frame, text="", font=("Arial", 9),
                                       fg="#7a9ab8", bg="#1e2a3a")
        self.progress_label.pack(anchor="w")
        self.progress_bar = ttk.Progressbar(prog_frame, mode="determinate")
        self.progress_bar.pack(fill="x", pady=(2, 4))
        self.cost_label = tk.Label(prog_frame, text="", font=("Arial", 9),
                                   fg="#ff9f4a", bg="#1e2a3a", anchor="w")
        self.cost_label.pack(anchor="w", pady=(0, 4))
        style = ttk.Style()
        style.theme_use("default")
        style.configure("TProgressbar", troughcolor="#2a3f55", background="#4da6ff", thickness=12)

        # Log
        log_outer = tk.Frame(self, bg="#1e2a3a")
        log_outer.pack(fill="both", expand=True, padx=20, pady=(0, 6))
        tk.Label(log_outer, text="Live Output", font=("Arial", 9, "bold"),
                 fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w")
        log_inner = tk.Frame(log_outer, bg="#0d1b2a")
        log_inner.pack(fill="both", expand=True)
        self.log = tk.Text(log_inner, bg="#0d1b2a", fg="#c8d8e8", font=("Consolas", 9),
                           relief="flat", wrap="word", state="disabled")
        sb = tk.Scrollbar(log_inner, command=self.log.yview, bg="#1e2a3a")
        self.log.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.log.pack(fill="both", expand=True, padx=6, pady=6)
        self.log.tag_config("info",    foreground="#4da6ff")
        self.log.tag_config("success", foreground="#4ddb8a")
        self.log.tag_config("warning", foreground="#f0c040")
        self.log.tag_config("error",   foreground="#ff6b6b")
        self.log.tag_config("dim",     foreground="#5a7a95")
        self.log.tag_config("cite",    foreground="#7a6a95")

        # Bottom buttons
        btn_frame = tk.Frame(self, bg="#1e2a3a", pady=10)
        btn_frame.pack(fill="x", padx=20)

        self.start_btn = tk.Button(btn_frame, text="▶  Start", font=("Arial", 11, "bold"),
                                   bg="#1a6b3a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._start)
        self.start_btn.pack(side="left", padx=(0, 6))
        self.stop_btn = tk.Button(btn_frame, text="⏹  Stop", font=("Arial", 11, "bold"),
                                  bg="#6b1a1a", fg="white", relief="flat", padx=20, pady=8,
                                  cursor="hand2", state="disabled", command=self._stop)
        self.stop_btn.pack(side="left", padx=(0, 6))
        self.pause_btn = tk.Button(btn_frame, text="⏸  Pause", font=("Arial", 11, "bold"),
                                   bg="#4a4a1a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._pause)
        self.pause_btn.pack(side="left", padx=(0, 6))
        self.retry_btn = tk.Button(btn_frame, text="🔁  Retry Failed", font=("Arial", 11),
                                   bg="#6b3a1a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._retry_failed)
        self.retry_btn.pack(side="left")
        self.export_btn = tk.Button(btn_frame, text="💾  Export Results", font=("Arial", 11),
                                    bg="#1a4a6b", fg="white", relief="flat", padx=20, pady=8,
                                    cursor="hand2", state="disabled", command=self._export)
        self.export_btn.pack(side="right")

    # ── API Key ───────────────────────────────────────────────────────────────

    def _test_api_key(self):
        key = self.api_key.get().strip()
        if not key:
            messagebox.showerror("No Key", "Enter a Perplexity API key first.")
            return
        self.key_status.config(text="Testing key...", fg="#f0c040")
        self.update()

        def _do_test():
            t0     = time.time()
            result = check_business(key, "Target", "https://www.target.com", "Minneapolis", "MN", max_retries=1)
            elapsed = time.time() - t0
            if result["error"]:
                self.after(0, lambda: self.key_status.config(
                    text=f"✗ Key test failed: {result['error'][:60]}", fg="#ff6b6b"))
            else:
                self.after(0, lambda: self.key_status.config(
                    text=f"✓ Key works! ({elapsed:.1f}s response)", fg="#4ddb8a"))

        threading.Thread(target=_do_test, daemon=True).start()

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _log(self, msg, tag="dim"):
        def _do():
            self.log.configure(state="normal")
            self._log_lines += 1
            if self._log_lines > LOG_MAX_LINES:
                self.log.delete("1.0", "2.0")
                self._log_lines -= 1
            self.log.insert("end", msg + "\n", tag)
            self.log.see("end")
            self.log.configure(state="disabled")
        self.after(0, _do)

    def _set_progress(self, value, label_text=None):
        def _do():
            self.progress_bar["value"] = value
            if label_text is not None:
                self.progress_label.config(text=label_text)
        self.after(0, _do)

    def _set_buttons(self, start=None, stop=None, export=None, pause=None, retry=None):
        def _do():
            if start  is not None: self.start_btn.config( state=start)
            if stop   is not None: self.stop_btn.config(  state=stop)
            if export is not None: self.export_btn.config(state=export)
            if pause  is not None: self.pause_btn.config( state=pause)
            if retry  is not None: self.retry_btn.config( state=retry)
        self.after(0, _do)

    def _eta_string(self, completed_so_far: int, total: int) -> str:
        """Estimate remaining time based on actual recent throughput."""
        if len(self._throughput) < 2:
            return ""
        elapsed = self._throughput[-1] - self._throughput[0]
        rate    = len(self._throughput) / elapsed if elapsed > 0 else 0  # rows/sec
        if rate == 0:
            return ""
        remaining = total - completed_so_far
        secs      = remaining / rate
        if secs < 60:
            return f"~{int(secs)}s remaining"
        elif secs < 3600:
            return f"~{int(secs/60)}m remaining"
        else:
            return f"~{secs/3600:.1f}h remaining"

    # ── File Handling ─────────────────────────────────────────────────────────

    def _load_and_confirm(self, path, already_done=0):
        try:
            wb      = openpyxl.load_workbook(path, read_only=True)
            ws      = wb.active
            headers = [cell.value for cell in next(ws.rows)]
            total   = sum(1 for _ in ws.rows) - 1
            wb.close()
        except Exception as e:
            messagebox.showerror("Error", f"Could not read Excel file:\n{e}")
            return None, None, None

        detected = detect_columns(headers)
        dialog   = ColumnConfirmDialog(self, headers, detected,
                                       total_rows=total, already_done=already_done,
                                       workers=self.num_workers)
        self.wait_window(dialog)
        if dialog.result is None:
            return None, None, None
        return dialog.result, dialog.chosen_workers, total

    def _new_run(self):
        path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel Files", "*.xlsx *.xls"), ("All Files", "*.*")]
        )
        if not path:
            return

        col_map, workers, total = self._load_and_confirm(path, already_done=0)
        if col_map is None:
            return

        self.col_map     = col_map
        self.num_workers = workers
        self._total_rows = total
        self._done_at_start = 0

        script_dir = os.path.dirname(os.path.abspath(__file__))
        file_stem  = os.path.splitext(os.path.basename(path))[0]
        timestamp  = datetime.now().strftime("%Y-%m-%d_%H%M")
        run_name   = f"{file_stem}_{timestamp}"
        run_dir    = os.path.join(script_dir, "Runs", run_name)
        os.makedirs(run_dir, exist_ok=True)

        dest = os.path.join(run_dir, os.path.basename(path))
        shutil.copy2(path, dest)

        self.run_dir = run_dir
        self.run_info.config(text=f"Run folder: Runs/{run_name}", fg="#4ddb8a")
        self._log(f"New run: {run_name}", "info")
        self._log(f"File: {os.path.basename(dest)}  |  {total} businesses  |  {workers} worker(s)", "dim")
        self._set_buttons(start="normal", export="disabled", retry="disabled")

    def _resume_run(self):
        folder = filedialog.askdirectory(title="Select a Run Folder to Resume")
        if not folder:
            return

        xlsx_files = [f for f in os.listdir(folder)
                      if f.endswith((".xlsx", ".xls")) and "Results" not in f]
        if not xlsx_files:
            messagebox.showerror("Error", "No input Excel file found in that folder.")
            return

        checkpoint_path = os.path.join(folder, "checkpoint.csv")
        done_pre        = load_checkpoint(checkpoint_path)
        xlsx_path       = os.path.join(folder, xlsx_files[0])

        col_map, workers, total = self._load_and_confirm(xlsx_path, already_done=len(done_pre))
        if col_map is None:
            return

        self.col_map        = col_map
        self.num_workers    = workers
        self.run_dir        = folder
        self._total_rows    = total or 0
        self._done_at_start = len(done_pre)

        self.run_info.config(
            text=f"Resuming: {os.path.basename(folder)}  ({len(done_pre)} rows already done)",
            fg="#f0c040"
        )
        self._log(f"Resuming: {os.path.basename(folder)}", "warning")
        self._log(f"Already done: {len(done_pre)}  |  Workers: {workers}", "dim")

        summary = get_run_summary(checkpoint_path)
        self._set_buttons(
            start="normal",
            export="normal" if done_pre else "disabled",
            retry="normal" if summary["Errors"] > 0 else "disabled",
        )

    # ── Run Control ───────────────────────────────────────────────────────────

    def _start(self):
        if not self.api_key.get().strip():
            messagebox.showerror("Missing API Key", "Please enter your Perplexity API key.")
            return
        if not self.run_dir:
            messagebox.showerror("No Run", "Please select New Run or Resume Run first.")
            return

        self.stop_flag        = False
        self.pause_flag       = False
        self.running          = True
        self._throughput      = []
        self._session_cost    = 0.0
        self._session_checked = 0
        self.after(0, lambda: self.cost_label.config(text=""))
        self._set_buttons(start="disabled", stop="normal", export="disabled",
                          pause="normal", retry="disabled")
        threading.Thread(target=self._run_worker, daemon=True).start()

    def _stop(self):
        self.stop_flag  = True
        self.pause_flag = False
        self._log("Stop requested — finishing in-flight checks...", "warning")
        self._set_buttons(stop="disabled", pause="disabled")

    def _pause(self):
        self.pause_flag = True
        self._log("Paused — click Resume to continue.", "warning")
        self.after(0, lambda: self.pause_btn.config(
            text="▶  Resume", bg="#1a6b3a", command=self._resume_pause))

    def _resume_pause(self):
        if self.stop_flag:
            return
        self.pause_flag = False
        self._log("Resumed.", "info")
        self.after(0, lambda: self.pause_btn.config(
            text="⏸  Pause", bg="#4a4a1a", command=self._pause))

    def _retry_failed(self):
        if not self.run_dir or self.running:
            messagebox.showwarning("Run Active", "Please stop the current run before retrying.")
            return

        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")
        done            = load_checkpoint(checkpoint_path)
        good_rows, retry_count = [], 0

        for r in done.values():
            is_error = (
                r.get("AI_Confidence", "1") == "0" and
                any(kw in r.get("AI_Evidence", "") for kw in ERROR_KEYWORDS)
            )
            if is_error:
                retry_count += 1
            else:
                good_rows.append(r)

        if retry_count == 0:
            messagebox.showinfo("No Failures", "No failed rows found to retry.")
            return

        if not messagebox.askyesno(
            "Retry Failed Rows",
            f"Found {retry_count} failed rows.\n\n"
            "They will be removed from the checkpoint and re-checked on the next run.\n\nContinue?"
        ):
            return

        rewrite_checkpoint(checkpoint_path, good_rows)
        self._log(f"Removed {retry_count} failed rows from checkpoint.", "warning")
        self._set_buttons(retry="disabled", start="normal")

    def _export(self):
        if not self.run_dir:
            messagebox.showerror("No Run", "No active run to export.")
            return

        xlsx_files = [f for f in os.listdir(self.run_dir)
                      if f.endswith((".xlsx", ".xls")) and "Results" not in f]
        if not xlsx_files:
            messagebox.showerror("Error", "No input Excel file found in run folder.")
            return

        input_path      = os.path.join(self.run_dir, xlsx_files[0])
        stem            = os.path.splitext(xlsx_files[0])[0]
        output_path     = os.path.join(self.run_dir, f"{stem}_Results.xlsx")
        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")

        try:
            saved = build_output_excel(input_path, output_path, checkpoint_path)
            self._log(f"Exported: {os.path.basename(saved)}", "success")
            if messagebox.askyesno("Exported", f"Results saved to:\n{saved}\n\nOpen the file now?"):
                open_file(saved)
        except PermissionError:
            messagebox.showerror("File In Use",
                                 "Could not save — the Results file may be open in Excel.\n"
                                 "Close it and try again.")
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

    def _on_close(self):
        if self.running:
            if not messagebox.askyesno(
                "Run In Progress",
                "A run is currently active. Progress is saved to checkpoint.\n\nQuit anyway?"
            ):
                return
            self.stop_flag  = True
            self.pause_flag = False
        self.destroy()

    # ── Background Worker ─────────────────────────────────────────────────────

    def _run_worker(self):
        api_key = self.api_key.get().strip()

        xlsx_files = [f for f in os.listdir(self.run_dir)
                      if f.endswith((".xlsx", ".xls")) and "Results" not in f]
        if not xlsx_files:
            self._log("No Excel file found in run folder.", "error")
            self._set_buttons(start="normal", stop="disabled", pause="disabled")
            self.running = False
            return

        input_path      = os.path.join(self.run_dir, xlsx_files[0])
        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")

        try:
            # Load workbook and snapshot all row data into plain dicts BEFORE spawning
            # threads. openpyxl worksheets are not thread-safe — ws.cell() mutates
            # internal state and will corrupt reads when called concurrently.
            wb       = openpyxl.load_workbook(input_path, read_only=True)
            ws_snap  = wb.active
            all_rows = list(ws_snap.rows)
            wb.close()

            total_rows = len(all_rows) - 1  # subtract header (max_row can be None)
            done       = load_checkpoint(checkpoint_path)
            col        = self.col_map

            def _val(row, col_idx):
                if col_idx is None or col_idx >= len(row):
                    return ""
                return str(row[col_idx].value or "").strip()

            # Plain dict cache — safe to read from multiple threads simultaneously
            row_cache = {}
            for i, row in enumerate(all_rows[1:], start=2):
                row_cache[i] = {
                    "name":    _val(row, col.get("name")),
                    "website": _val(row, col.get("website")),
                    "city":    _val(row, col.get("city")),
                    "state":   _val(row, col.get("state")),
                }

            self._log(f"File: {os.path.basename(input_path)}", "info")
            self._log(f"Total: {total_rows}  |  Done: {len(done)}  |  "
                      f"Remaining: {total_rows - len(done)}  |  Workers: {self.num_workers}", "info")
            self._log("─" * 55, "dim")

            self._total_rows    = total_rows
            self._done_at_start = len(done)
            self.after(0, lambda: self.progress_bar.config(maximum=total_rows, value=len(done)))

            # Build list of pending rows
            pending = [row_idx for row_idx in row_cache if str(row_idx) not in done]

            # Thread-safe counter
            counter_lock    = threading.Lock()
            completed_count = [len(done)]

            def process_row(row_idx: int):
                # Check pause/stop BEFORE starting the API call — this is where
                # pause actually takes effect (not in the submission loop below).
                while self.pause_flag and not self.stop_flag:
                    time.sleep(0.3)
                if self.stop_flag:
                    return

                d       = row_cache[row_idx]
                name    = d["name"]
                website = d["website"]
                city    = d["city"]
                state   = d["state"]

                result     = check_business(api_key, name, website, city, state)
                checked_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

                save_checkpoint(checkpoint_path, row_idx, name, website,
                                result["status"], result["confidence"],
                                result["evidence"], checked_at,
                                requires_review=bool(result.get("requires_review", False)),
                                review_reason=str(result.get("review_reason") or ""))

                with counter_lock:
                    completed_count[0] += 1
                    n = completed_count[0]
                    self._throughput.append(time.time())
                    if len(self._throughput) > 50:
                        self._throughput = self._throughput[-50:]
                    self._session_cost    += result.get("cost_usd", 0.0)
                    self._session_checked += 1
                    session_cost    = self._session_cost
                    session_checked = self._session_checked

                progress = n
                eta      = self._eta_string(n, total_rows)
                label    = f"Checking {progress} of {total_rows}: {name}{('  ' + eta) if eta else ''}"
                self._set_progress(progress, label)

                avg = session_cost / session_checked if session_checked > 0 else 0
                self.after(0, lambda sc=session_cost, av=avg: self.cost_label.config(
                    text=f"API cost this session: ${sc:.4f}  (avg ${av:.4f}/check)"))

                tag = ("success" if result["status"] == "Active" else
                       "error"   if result["status"] == "Likely Closed" else "warning")
                self._log(f"[{progress}/{total_rows}] {name[:35]} ({city}, {state})", "dim")
                self._log(f"  → {result['status']} ({result['confidence']}%) | {result['evidence'][:100]}", tag)
                if result.get("citations"):
                    self._log(f"     {' | '.join(result['citations'][:2])}", "cite")

            with ThreadPoolExecutor(max_workers=self.num_workers) as executor:
                futures = {
                    executor.submit(process_row, row_idx): row_idx
                    for row_idx in pending
                }
                for future in as_completed(futures):
                    if self.stop_flag:
                        break
                    try:
                        future.result()
                    except Exception as e:
                        self._log(f"Worker error: {e}", "error")

            if self.stop_flag:
                self._log("Stopped. Progress saved to checkpoint.", "warning")
            else:
                self._log("─" * 55, "dim")
                self._log("All businesses checked!", "success")
                self._set_progress(total_rows, "Complete!")
                summary = get_run_summary(checkpoint_path)
                self.after(0, lambda: SummaryDialog(
                    self, summary, on_export=self._export, on_retry=self._retry_failed))
                if summary["Errors"] > 0:
                    self._set_buttons(retry="normal")

        except Exception as e:
            self._log(f"Fatal error: {e}", "error")
            err = str(e)
            self.after(0, lambda: messagebox.showerror("Error", err))

        finally:
            self.running = False
            self._set_buttons(start="normal", stop="disabled", pause="disabled", export="normal")
            self.after(0, lambda: self.pause_btn.config(
                text="⏸  Pause", bg="#4a4a1a", command=self._pause))


if __name__ == "__main__":
    app = App()
    app.mainloop()
