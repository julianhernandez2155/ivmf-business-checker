"""
Excel output builder.

Reads the original input spreadsheet and the checkpoint CSV,
then produces a new Excel file with AI result columns first (left side)
followed by the original source columns.
"""

import re

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter
from tools.checkpoint import load_checkpoint

# Characters illegal in XML 1.0 — strip them before writing cells
_ILLEGAL_XML_RE = re.compile(
    r"[\x00-\x08\x0B\x0C\x0E-\x1F\x7F\uFFFE\uFFFF]"
)


def _safe(value):
    """Strip XML-illegal characters from string values."""
    if isinstance(value, str):
        return _ILLEGAL_XML_RE.sub("", value)
    return value

# AI result columns — placed first so analysts see outcomes immediately on open
_AI_COLS = ["AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"]

# Status cell (AI_Status column) fill + font
_STATUS_CELL_FILLS = {
    "Active":          PatternFill("solid", fgColor="C6EFCE"),
    "Likely Closed":   PatternFill("solid", fgColor="FFC7CE"),
    "Uncertain":       PatternFill("solid", fgColor="FFEB9C"),
    "No Web Presence": PatternFill("solid", fgColor="DDDDDD"),
}
_STATUS_CELL_FONTS = {
    "Active":          Font(color="276221", bold=True),
    "Likely Closed":   Font(color="9C0006", bold=True),
    "Uncertain":       Font(color="9C5700", bold=True),
    "No Web Presence": Font(color="595959", bold=True),
}

# Full-row background fill (lighter shade of each status color)
_STATUS_ROW_FILLS = {
    "Active":          PatternFill("solid", fgColor="E2EFDA"),
    "Likely Closed":   PatternFill("solid", fgColor="FCE4D6"),
    "Uncertain":       PatternFill("solid", fgColor="FFF2CC"),
    "No Web Presence": PatternFill("solid", fgColor="F2F2F2"),
}

# Column widths for the AI columns (characters)
_AI_COL_WIDTHS = {
    "AI_Status":     16,
    "AI_Confidence": 14,
    "AI_Evidence":   70,
    "AI_Checked_At": 20,
}

# Width hints for source columns — matched by substring against the header name (lowercase)
_SOURCE_COL_WIDTH_HINTS = [
    (["business name", "company name", "organization name", "name"], 32),
    (["website", "web site", "url", "site", "link"],                 35),
    (["email", "e-mail", "e mail"],                                  30),
    (["phone", "telephone", "tel ", "mobile", "cell"],               18),
    (["city", "town", "municipality"],                               16),
    (["state", "province", "region"],                                 8),
    (["zip", "postal", "post code"],                                 10),
    (["address", "street", "suite"],                                 30),
    (["owner", "contact", "principal"],                              22),
    (["description", "notes", "comment", "detail"],                 30),
]
_DEFAULT_SOURCE_COL_WIDTH = 18

# Header row style
_HEADER_FILL = PatternFill("solid", fgColor="2F5496")
_HEADER_FONT = Font(color="FFFFFF", bold=True)


def _source_col_width(header: str) -> int:
    """Return a width for a source column based on its header name."""
    h = str(header or "").lower().strip()
    for hints, width in _SOURCE_COL_WIDTH_HINTS:
        if any(hint in h for hint in hints):
            return width
    return _DEFAULT_SOURCE_COL_WIDTH


def _is_name_col(header: str) -> bool:
    h = str(header or "").lower().strip()
    return any(k in h for k in ["business name", "company name", "organization name", "name"])


def build_output_excel(input_path: str, output_path: str, checkpoint_path: str) -> str:
    """
    Build the results Excel file.

    Column order: AI_Status | AI_Confidence | AI_Evidence | AI_Checked_At | [original columns]

    Formatting:
    - AI result columns are first so outcomes are visible without scrolling
    - Freeze panes at E2 (header row + AI columns stay locked)
    - Auto-filter on all columns
    - Full-row background color per status
    - AI_Status cell gets stronger fill + bold font
    - Column widths set for readability

    Args:
        input_path:      Path to the original input .xlsx file.
        output_path:     Where to save the results .xlsx file.
        checkpoint_path: Path to the checkpoint.csv file.

    Returns:
        output_path on success.
    """
    checkpoint = load_checkpoint(checkpoint_path)

    # Use read_only mode for memory efficiency on large files.
    # NOTE: read_only workbooks must be explicitly closed.
    wb_in = openpyxl.load_workbook(input_path, read_only=True)
    ws_in = wb_in.active

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Results"

    # Header row: AI columns first, then original source columns
    source_headers = [cell.value for cell in next(ws_in.rows)]
    headers = _AI_COLS + source_headers
    ai_col_count = len(_AI_COLS)

    # Pre-compute which source columns are "name" columns (for bold)
    name_col_indices = {
        ai_col_count + i + 1   # 1-based output column index
        for i, h in enumerate(source_headers)
        if _is_name_col(h)
    }

    ws_out.append(headers)

    # Style the header row
    ws_out.row_dimensions[1].height = 22
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws_out.cell(row=1, column=col_idx)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Write each data row
    for row_idx, row in enumerate(ws_in.rows, start=1):
        if row_idx == 1:
            continue  # skip header (already written above)

        source_data = [_safe(cell.value) for cell in row]
        key = str(row_idx)

        if key in checkpoint:
            r = checkpoint[key]
            status = r["AI_Status"]
            ai_data = [
                status,
                r["AI_Confidence"],
                _safe(r["AI_Evidence"]),
                r["AI_Checked_At"],
            ]
        else:
            status = None
            ai_data = ["Not Checked", "", "", ""]

        ws_out.append(ai_data + source_data)
        out_row = ws_out.max_row

        if status in _STATUS_ROW_FILLS:
            row_fill = _STATUS_ROW_FILLS[status]
            total_cols = len(headers)
            for col_idx in range(1, total_cols + 1):
                ws_out.cell(row=out_row, column=col_idx).fill = row_fill

            # AI_Status cell (column 1) gets stronger fill + bold font
            status_cell = ws_out.cell(row=out_row, column=1)
            status_cell.fill = _STATUS_CELL_FILLS[status]
            status_cell.font = _STATUS_CELL_FONTS[status]
            status_cell.alignment = Alignment(horizontal="center")

        # Bold the business name column(s)
        for col_idx in name_col_indices:
            ws_out.cell(row=out_row, column=col_idx).font = Font(bold=True)

    wb_in.close()  # required when read_only=True

    # Freeze panes: lock header row + AI columns so source data scrolls freely
    ws_out.freeze_panes = "E2"

    # Auto-filter across all columns
    ws_out.auto_filter.ref = ws_out.dimensions

    # AI column widths + wrap for evidence
    for col_name, width in _AI_COL_WIDTHS.items():
        col_idx = _AI_COLS.index(col_name) + 1
        col_letter = get_column_letter(col_idx)
        ws_out.column_dimensions[col_letter].width = width
    # Wrap AI_Evidence so long text stays readable without expanding the row height excessively
    evidence_col = get_column_letter(_AI_COLS.index("AI_Evidence") + 1)
    for row_idx in range(2, ws_out.max_row + 1):
        ws_out[f"{evidence_col}{row_idx}"].alignment = Alignment(wrap_text=True, vertical="top")

    # Source column widths — detected by header name
    for i, header in enumerate(source_headers):
        col_idx = ai_col_count + i + 1
        ws_out.column_dimensions[get_column_letter(col_idx)].width = _source_col_width(header)

    wb_out.save(output_path)
    return output_path
