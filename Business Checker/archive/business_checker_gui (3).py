"""
Business Operational Status Checker — GUI Version
Double-click this file to launch. No terminal needed.

SETUP (one time):
    pip install anthropic openpyxl
"""

import os
import csv
import time
import json
import shutil
import threading
import openpyxl
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import anthropic

# ── Constants ─────────────────────────────────────────────────────────────────

DELAY_SECONDS  = 8
MAX_RETRIES    = 4
MODEL          = "claude-haiku-4-5-20251001"
LOG_MAX_LINES  = 500
API_KEY_FILE   = os.path.join(os.path.expanduser("~"), ".biz_checker_key")

HEADER_HINTS = {
    "name":    ["name", "business name", "company", "organization"],
    "website": ["website", "web", "url", "site", "link"],
    "city":    ["city", "town", "municipality"],
    "state":   ["state", "province", "region"],
}

# ── API Key Persistence ───────────────────────────────────────────────────────

def load_saved_api_key():
    """Load API key saved from a previous session."""
    try:
        if os.path.exists(API_KEY_FILE):
            with open(API_KEY_FILE, "r") as f:
                return f.read().strip()
    except Exception:
        pass
    return ""

def save_api_key(key):
    """Save API key to user home directory for next session."""
    try:
        with open(API_KEY_FILE, "w") as f:
            f.write(key.strip())
    except Exception:
        pass

# ── Core Logic ────────────────────────────────────────────────────────────────

def detect_columns(headers):
    detected = {"name": None, "website": None, "city": None, "state": None}
    for idx, h in enumerate(headers):
        if h is None:
            continue
        h_lower = str(h).lower().strip()
        for field, hints in HEADER_HINTS.items():
            if detected[field] is None:
                if any(hint in h_lower for hint in hints):
                    detected[field] = idx
    return detected


def load_checkpoint(path):
    done = {}
    if not os.path.exists(path):
        return done
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            done[row["row_index"]] = row
    return done


def save_checkpoint(path, row_index, name, website, status, confidence, evidence, checked_at):
    file_exists = os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "row_index", "name", "website",
            "AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"
        ])
        if not file_exists:
            writer.writeheader()
        writer.writerow({
            "row_index":     row_index,
            "name":          name,
            "website":       website,
            "AI_Status":     status,
            "AI_Confidence": confidence,
            "AI_Evidence":   evidence,
            "AI_Checked_At": checked_at,
        })


def rewrite_checkpoint(path, rows):
    """Overwrite checkpoint with a new set of rows (used by retry failed)."""
    with open(path, "w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "row_index", "name", "website",
            "AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"
        ])
        writer.writeheader()
        for r in rows:
            writer.writerow(r)


def get_run_summary(checkpoint_path):
    """Return a dict of status counts from a checkpoint file."""
    done = load_checkpoint(checkpoint_path)
    summary = {"Active": 0, "Likely Closed": 0, "Uncertain": 0,
               "No Web Presence": 0, "Errors": 0, "Total": len(done)}
    for r in done.values():
        s = r.get("AI_Status", "")
        conf = r.get("AI_Confidence", "1")
        is_error = (conf == "0" and any(
            kw in r.get("AI_Evidence", "")
            for kw in ["Parse error", "No text response", "Error:", "Failed after"]
        ))
        if is_error:
            summary["Errors"] += 1
        elif s in summary:
            summary[s] += 1
        else:
            summary["Uncertain"] += 1
    return summary


def check_business(client, name, website, city, state):
    location     = ", ".join(filter(None, [city, state]))
    website_note = f"Their listed website is: {website}" if website else "No website listed."

    prompt = f"""Search the web for "{name}" located in {location} and determine if this business is still actively operating.

Check their website ({website_note}), Google, Yelp, or any other sources for signs of active operation or closure.

After searching, you MUST respond with ONLY this JSON and nothing else — no intro, no explanation, no markdown:
{{"status": "Active", "confidence": 85, "evidence": "Website is live with recent activity."}}

Status must be one of: Active, Likely Closed, Uncertain, No Web Presence
Confidence must be an integer 0-100."""

    for attempt in range(MAX_RETRIES):
        try:
            messages = [{"role": "user", "content": prompt}]
            response = client.messages.create(
                model=MODEL,
                max_tokens=200,
                tools=[{"type": "web_search_20250305", "name": "web_search"}],
                messages=messages
            )

            if response.stop_reason == "tool_use":
                messages.append({"role": "assistant", "content": response.content})
                messages.append({"role": "user", "content": "Now respond with ONLY the JSON result. No explanation, no markdown, just the JSON object."})
                response = client.messages.create(
                    model=MODEL,
                    max_tokens=200,
                    tools=[{"type": "web_search_20250305", "name": "web_search"}],
                    messages=messages
                )

            text_blocks = []
            for block in response.content:
                if hasattr(block, "type") and block.type == "text":
                    text_blocks.append(block.text.strip())

            if not text_blocks:
                time.sleep(5)
                continue

            result_text = text_blocks[-1]
            result_text = result_text.replace("```json", "").replace("```", "").strip()
            start = result_text.find("{")
            end   = result_text.rfind("}") + 1
            if start != -1 and end > start:
                result_text = result_text[start:end]

            parsed = json.loads(result_text)
            return (
                parsed.get("status", "Uncertain"),
                str(parsed.get("confidence", 50)),
                parsed.get("evidence", "No evidence found.")
            )

        except json.JSONDecodeError:
            return ("Uncertain", "0", f"Parse error: {result_text[:100]}")

        except anthropic.RateLimitError:
            wait = 30 * (attempt + 1)
            time.sleep(wait)
            continue

        except Exception as e:
            return ("Uncertain", "0", f"Error: {str(e)[:100]}")

    return ("Uncertain", "0", "Failed after max retries.")


def build_output_excel(input_path, output_path, checkpoint_path):
    checkpoint = load_checkpoint(checkpoint_path)
    wb_in  = openpyxl.load_workbook(input_path)
    ws_in  = wb_in.active
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Results"

    headers = [ws_in.cell(1, c).value for c in range(1, ws_in.max_column + 1)]
    headers += ["AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"]
    ws_out.append(headers)

    for row_idx in range(2, ws_in.max_row + 1):
        row_data = [ws_in.cell(row_idx, c).value for c in range(1, ws_in.max_column + 1)]
        key = str(row_idx)
        if key in checkpoint:
            r = checkpoint[key]
            row_data += [r["AI_Status"], r["AI_Confidence"], r["AI_Evidence"], r["AI_Checked_At"]]
        else:
            row_data += ["Not Checked", "", "", ""]
        ws_out.append(row_data)

    wb_out.save(output_path)
    return output_path


def get_cell_value(ws, row_idx, col_idx):
    if col_idx is None:
        return ""
    return ws.cell(row_idx, col_idx + 1).value or ""


# ── Summary Dialog ────────────────────────────────────────────────────────────

class SummaryDialog(tk.Toplevel):
    """Shows a run summary with stats and options to export or retry failed rows."""

    def __init__(self, parent, summary, on_export, on_retry):
        super().__init__(parent)
        self.title("Run Complete — Summary")
        self.geometry("420x360")
        self.configure(bg="#1e2a3a")
        self.resizable(False, False)
        self.grab_set()

        total   = summary["Total"]
        active  = summary["Active"]
        closed  = summary["Likely Closed"]
        uncert  = summary["Uncertain"]
        no_web  = summary["No Web Presence"]
        errors  = summary["Errors"]

        tk.Label(self, text="Run Complete!", font=("Arial", 14, "bold"),
                 fg="#4ddb8a", bg="#1e2a3a").pack(pady=(18, 4))
        tk.Label(self, text="Here's a breakdown of results:",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1e2a3a").pack()

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=12)

        stats_frame = tk.Frame(self, bg="#0d1b2a", padx=20, pady=14)
        stats_frame.pack(fill="x", padx=20)

        def pct(n):
            return f"({int(n/total*100)}%)" if total > 0 else "(0%)"

        rows = [
            ("✓  Active",           str(active),  "#4ddb8a", pct(active)),
            ("✗  Likely Closed",    str(closed),  "#ff6b6b", pct(closed)),
            ("?  Uncertain",        str(uncert),  "#f0c040", pct(uncert)),
            ("○  No Web Presence",  str(no_web),  "#a0b8cc", pct(no_web)),
            ("⚠  Errors / Retries", str(errors),  "#ff9f4a", pct(errors)),
            ("─────────────────", "", "#2e4057", ""),
            ("Total Checked",       str(total),   "#c8d8e8", ""),
        ]

        for i, (label, val, color, pct_str) in enumerate(rows):
            tk.Label(stats_frame, text=label, font=("Arial", 10),
                     fg=color, bg="#0d1b2a", anchor="w", width=22).grid(
                         row=i, column=0, sticky="w", pady=2)
            tk.Label(stats_frame, text=val, font=("Arial", 10, "bold"),
                     fg=color, bg="#0d1b2a", anchor="e", width=6).grid(
                         row=i, column=1, sticky="e", pady=2)
            tk.Label(stats_frame, text=pct_str, font=("Arial", 9),
                     fg="#5a7a95", bg="#0d1b2a", anchor="w", width=8).grid(
                         row=i, column=2, sticky="w", padx=(6, 0), pady=2)

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=12)

        btn_frame = tk.Frame(self, bg="#1e2a3a")
        btn_frame.pack(pady=(0, 16))

        tk.Button(btn_frame, text="💾  Export Results", font=("Arial", 10, "bold"),
                  bg="#1a4a6b", fg="white", relief="flat", padx=14, pady=6,
                  cursor="hand2", command=lambda: [self.destroy(), on_export()]
                  ).pack(side="left", padx=6)

        retry_btn = tk.Button(btn_frame, text=f"🔁  Retry {errors} Failed", font=("Arial", 10),
                              bg="#6b3a1a", fg="white", relief="flat", padx=14, pady=6,
                              cursor="hand2", command=lambda: [self.destroy(), on_retry()])
        retry_btn.pack(side="left", padx=6)
        if errors == 0:
            retry_btn.config(state="disabled", bg="#2e4057", fg="#5a7a95")

        tk.Button(btn_frame, text="Close", font=("Arial", 10),
                  bg="#2e4057", fg="#a0b8cc", relief="flat", padx=14, pady=6,
                  cursor="hand2", command=self.destroy).pack(side="left", padx=6)


# ── Column Confirmation Dialog ────────────────────────────────────────────────

class ColumnConfirmDialog(tk.Toplevel):
    def __init__(self, parent, headers, detected, total_rows=0, already_done=0):
        super().__init__(parent)
        self.title("Confirm Column Mapping")
        self.geometry("520x440")
        self.configure(bg="#1e2a3a")
        self.resizable(False, False)
        self.grab_set()

        self.headers        = headers
        self.result         = None
        self.chosen_delay   = DELAY_SECONDS
        self.vars           = {}
        self.selected_delay = tk.IntVar(value=DELAY_SECONDS)

        tk.Label(self, text="Confirm Column Mapping",
                 font=("Arial", 13, "bold"), fg="#4da6ff", bg="#1e2a3a").pack(
                     pady=(16, 4), padx=20, anchor="w")
        tk.Label(self, text="The program detected the following columns. Correct any mistakes before starting.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1e2a3a",
                 wraplength=480, justify="left").pack(padx=20, anchor="w")

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        grid = tk.Frame(self, bg="#1e2a3a")
        grid.pack(padx=24, fill="x")

        header_options     = [f"{i+1}: {h}" for i, h in enumerate(headers) if h is not None]
        header_options_map = {f"{i+1}: {h}": i for i, h in enumerate(headers) if h is not None}
        self.header_options_map = header_options_map

        fields = [
            ("Business Name column:", "name"),
            ("Website column:",       "website"),
            ("City column:",          "city"),
            ("State column:",         "state"),
        ]

        for row_i, (label, field) in enumerate(fields):
            tk.Label(grid, text=label, font=("Arial", 10), fg="#a0b8cc",
                     bg="#1e2a3a", anchor="w", width=24).grid(
                         row=row_i, column=0, pady=6, sticky="w")
            var     = tk.StringVar()
            det_idx = detected.get(field)
            if det_idx is not None and det_idx < len(headers) and headers[det_idx] is not None:
                var.set(f"{det_idx+1}: {headers[det_idx]}")
            elif header_options:
                var.set(header_options[0])
            cb = ttk.Combobox(grid, textvariable=var, values=header_options,
                              state="readonly", width=30, font=("Arial", 10))
            cb.grid(row=row_i, column=1, pady=6, padx=(8, 0), sticky="w")
            self.vars[field] = var

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        # ── Cost Estimate Panel ───────────────────────────────────────────
        remaining    = max(0, total_rows - already_done)
        cost_per_biz = 0.05

        est_frame = tk.Frame(self, bg="#0d1b2a", padx=14, pady=10)
        est_frame.pack(fill="x", padx=20, pady=(0, 6))

        tk.Label(est_frame, text="Run Estimate", font=("Arial", 9, "bold"),
                 fg="#4da6ff", bg="#0d1b2a").grid(
                     row=0, column=0, columnspan=6, sticky="w", pady=(0, 6))

        def stat(parent, label, value, color, row, col):
            tk.Label(parent, text=label, font=("Arial", 8),
                     fg="#5a7a95", bg="#0d1b2a").grid(
                         row=row, column=col*2, sticky="w", padx=(0, 4))
            tk.Label(parent, text=value, font=("Arial", 10, "bold"),
                     fg=color, bg="#0d1b2a").grid(
                         row=row, column=col*2+1, sticky="w", padx=(0, 20))

        stat(est_frame, "Total rows:",   str(total_rows),   "#c8d8e8", 1, 0)
        stat(est_frame, "Already done:", str(already_done), "#4ddb8a", 1, 1)
        stat(est_frame, "Remaining:",    str(remaining),    "#f0c040", 1, 2)

        tk.Label(est_frame, text="Est. cost:",  font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a").grid(row=2, column=0, sticky="w", padx=(0, 4))
        tk.Label(est_frame, text="Est. time:",  font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a").grid(row=2, column=2, sticky="w", padx=(0, 4))
        tk.Label(est_frame, text="Delay/call:", font=("Arial", 8), fg="#5a7a95", bg="#0d1b2a").grid(row=2, column=4, sticky="w", padx=(0, 4))

        cost_label  = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#ff9f4a", bg="#0d1b2a")
        time_label  = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#4da6ff", bg="#0d1b2a")
        delay_label = tk.Label(est_frame, font=("Arial", 10, "bold"), fg="#c8d8e8", bg="#0d1b2a")
        cost_label.grid( row=2, column=1, sticky="w", padx=(0, 20))
        time_label.grid( row=2, column=3, sticky="w", padx=(0, 20))
        delay_label.grid(row=2, column=5, sticky="w")

        def update_estimates():
            d         = self.selected_delay.get()
            est_secs  = remaining * d
            est_hours = est_secs / 3600
            est_time  = f"{est_hours:.1f} hrs" if est_hours >= 1 else f"{int(est_secs / 60)} min"
            cost_label.config( text=f"~${remaining * cost_per_biz:.2f}")
            time_label.config( text=est_time)
            delay_label.config(text=f"{d}s")

        update_estimates()

        def toggle_delay():
            new_delay = 4 if self.selected_delay.get() == 8 else 8
            self.selected_delay.set(new_delay)
            update_estimates()
            if new_delay == 4:
                toggle_btn.config(text="Switch to 8s (safer, slower)")
            else:
                toggle_btn.config(text="Switch to 4s (faster, slightly riskier)")

        toggle_btn = tk.Button(
            est_frame,
            text="Switch to 4s (faster, slightly riskier)",
            font=("Arial", 8), bg="#2e4057", fg="#a0b8cc",
            relief="flat", padx=8, pady=2, cursor="hand2",
            command=toggle_delay
        )
        toggle_btn.grid(row=3, column=0, columnspan=6, sticky="w", pady=(8, 0))

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=10)

        btn_frame = tk.Frame(self, bg="#1e2a3a")
        btn_frame.pack(pady=(0, 16))

        tk.Button(btn_frame, text="✓  Confirm Columns", font=("Arial", 10, "bold"),
                  bg="#1a6b3a", fg="white", relief="flat", padx=16, pady=6,
                  cursor="hand2", command=self._confirm).pack(side="left", padx=8)
        tk.Button(btn_frame, text="Cancel", font=("Arial", 10),
                  bg="#2e4057", fg="#a0b8cc", relief="flat", padx=16, pady=6,
                  cursor="hand2", command=self.destroy).pack(side="left", padx=8)

    def _confirm(self):
        mapping = {}
        for field, var in self.vars.items():
            val = var.get()
            if val in self.header_options_map:
                mapping[field] = self.header_options_map[val]
            else:
                messagebox.showerror("Error", f"Please select a valid column for: {field}")
                return
        self.result       = mapping
        self.chosen_delay = self.selected_delay.get()
        self.destroy()


# ── Main App ──────────────────────────────────────────────────────────────────

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Business Operational Status Checker")
        self.geometry("920x660")
        self.resizable(True, True)
        self.configure(bg="#1e2a3a")

        self.api_key         = tk.StringVar(value=load_saved_api_key())
        self.run_dir         = None
        self.col_map         = None
        self.chosen_delay    = DELAY_SECONDS
        self.stop_flag       = False
        self.pause_flag      = False
        self.running         = False  # True while worker thread is active
        self._log_line_count = 0

        self._build_ui()

    def _build_ui(self):
        # ── Header ────────────────────────────────────────────────────────
        hdr = tk.Frame(self, bg="#1e2a3a", pady=14)
        hdr.pack(fill="x", padx=20)
        tk.Label(hdr, text="Business Operational Status Checker",
                 font=("Arial", 16, "bold"), fg="#4da6ff", bg="#1e2a3a").pack(anchor="w")
        tk.Label(hdr, text="IVMF  |  Powered by Claude AI + Web Search",
                 font=("Arial", 10), fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w")
        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20)

        # ── API Key ───────────────────────────────────────────────────────
        api_frame = tk.Frame(self, bg="#1e2a3a", pady=12)
        api_frame.pack(fill="x", padx=20)
        tk.Label(api_frame, text="Anthropic API Key:", font=("Arial", 10),
                 fg="#a0b8cc", bg="#1e2a3a", width=18, anchor="w").grid(row=0, column=0, sticky="w")
        self.api_entry = tk.Entry(api_frame, textvariable=self.api_key, show="*",
                                  font=("Arial", 10), bg="#2a3f55", fg="white",
                                  insertbackground="white", relief="flat", width=50)
        self.api_entry.grid(row=0, column=1, sticky="ew", padx=(0, 8), ipady=5)
        tk.Button(api_frame, text="Show", font=("Arial", 9), bg="#2e4057", fg="#a0b8cc",
                  relief="flat", cursor="hand2",
                  command=lambda: self.api_entry.config(
                      show="" if self.api_entry.cget("show") == "*" else "*")
                  ).grid(row=0, column=2, padx=(0, 4))
        tk.Button(api_frame, text="Save", font=("Arial", 9), bg="#1a4a6b", fg="white",
                  relief="flat", cursor="hand2",
                  command=self._save_api_key
                  ).grid(row=0, column=3)
        api_frame.columnconfigure(1, weight=1)

        # Show saved key indicator
        self.key_status = tk.Label(api_frame,
                                   text="✓ Key loaded from last session" if load_saved_api_key() else "",
                                   font=("Arial", 8), fg="#4ddb8a", bg="#1e2a3a")
        self.key_status.grid(row=1, column=1, sticky="w", pady=(2, 0))

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20)

        # ── Run Mode Cards ────────────────────────────────────────────────
        run_frame = tk.Frame(self, bg="#1e2a3a", pady=14)
        run_frame.pack(fill="x", padx=20)
        tk.Label(run_frame, text="Select a run mode to begin:",
                 font=("Arial", 10), fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w", pady=(0, 8))

        modes = tk.Frame(run_frame, bg="#1e2a3a")
        modes.pack(fill="x")

        new_card = tk.Frame(modes, bg="#1a3a52", padx=16, pady=12, relief="flat")
        new_card.pack(side="left", fill="both", expand=True, padx=(0, 8))
        tk.Label(new_card, text="▶  New Run", font=("Arial", 11, "bold"),
                 fg="#4da6ff", bg="#1a3a52").pack(anchor="w")
        tk.Label(new_card, text="Pick any Excel file to start\na fresh check from scratch.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1a3a52", justify="left").pack(anchor="w", pady=(4, 8))
        tk.Button(new_card, text="Browse for Excel file", font=("Arial", 10),
                  bg="#1a6b3a", fg="white", relief="flat", padx=12, pady=6,
                  cursor="hand2", command=self._new_run).pack(anchor="w")

        res_card = tk.Frame(modes, bg="#1a3a52", padx=16, pady=12, relief="flat")
        res_card.pack(side="left", fill="both", expand=True)
        tk.Label(res_card, text="⏩  Resume Run", font=("Arial", 11, "bold"),
                 fg="#f0c040", bg="#1a3a52").pack(anchor="w")
        tk.Label(res_card, text="Pick a previous run folder inside\nthe 'Runs' folder to continue.",
                 font=("Arial", 9), fg="#7a9ab8", bg="#1a3a52", justify="left").pack(anchor="w", pady=(4, 8))
        tk.Button(res_card, text="Browse for run folder", font=("Arial", 10),
                  bg="#6b5a1a", fg="white", relief="flat", padx=12, pady=6,
                  cursor="hand2", command=self._resume_run).pack(anchor="w")

        tk.Frame(self, bg="#2e4057", height=1).pack(fill="x", padx=20, pady=(10, 0))

        # ── Run Info + Progress ───────────────────────────────────────────
        self.run_info = tk.Label(self, text="No run loaded.",
                                 font=("Arial", 9, "italic"), fg="#4a6a85", bg="#1e2a3a")
        self.run_info.pack(anchor="w", padx=22, pady=(6, 0))

        prog_frame = tk.Frame(self, bg="#1e2a3a")
        prog_frame.pack(fill="x", padx=20, pady=(4, 0))
        self.progress_label = tk.Label(prog_frame, text="",
                                       font=("Arial", 9), fg="#7a9ab8", bg="#1e2a3a")
        self.progress_label.pack(anchor="w")
        self.progress_bar = ttk.Progressbar(prog_frame, mode="determinate")
        self.progress_bar.pack(fill="x", pady=(2, 6))
        style = ttk.Style()
        style.theme_use("default")
        style.configure("TProgressbar", troughcolor="#2a3f55", background="#4da6ff", thickness=12)

        # ── Log ───────────────────────────────────────────────────────────
        log_outer = tk.Frame(self, bg="#1e2a3a")
        log_outer.pack(fill="both", expand=True, padx=20, pady=(0, 6))
        tk.Label(log_outer, text="Live Output", font=("Arial", 9, "bold"),
                 fg="#7a9ab8", bg="#1e2a3a").pack(anchor="w")
        log_inner = tk.Frame(log_outer, bg="#0d1b2a")
        log_inner.pack(fill="both", expand=True)
        self.log = tk.Text(log_inner, bg="#0d1b2a", fg="#c8d8e8", font=("Consolas", 9),
                           relief="flat", wrap="word", state="disabled")
        sb = tk.Scrollbar(log_inner, command=self.log.yview, bg="#1e2a3a")
        self.log.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.log.pack(fill="both", expand=True, padx=6, pady=6)
        self.log.tag_config("info",    foreground="#4da6ff")
        self.log.tag_config("success", foreground="#4ddb8a")
        self.log.tag_config("warning", foreground="#f0c040")
        self.log.tag_config("error",   foreground="#ff6b6b")
        self.log.tag_config("dim",     foreground="#5a7a95")

        # ── Bottom Buttons ────────────────────────────────────────────────
        btn_frame = tk.Frame(self, bg="#1e2a3a", pady=10)
        btn_frame.pack(fill="x", padx=20)

        self.start_btn = tk.Button(btn_frame, text="▶  Start", font=("Arial", 11, "bold"),
                                   bg="#1a6b3a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._start)
        self.start_btn.pack(side="left", padx=(0, 6))

        self.stop_btn = tk.Button(btn_frame, text="⏹  Stop", font=("Arial", 11, "bold"),
                                  bg="#6b1a1a", fg="white", relief="flat", padx=20, pady=8,
                                  cursor="hand2", state="disabled", command=self._stop)
        self.stop_btn.pack(side="left", padx=(0, 6))

        self.pause_btn = tk.Button(btn_frame, text="⏸  Pause", font=("Arial", 11, "bold"),
                                   bg="#4a4a1a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._pause)
        self.pause_btn.pack(side="left", padx=(0, 6))

        self.retry_btn = tk.Button(btn_frame, text="🔁  Retry Failed", font=("Arial", 11),
                                   bg="#6b3a1a", fg="white", relief="flat", padx=20, pady=8,
                                   cursor="hand2", state="disabled", command=self._retry_failed)
        self.retry_btn.pack(side="left")

        self.export_btn = tk.Button(btn_frame, text="💾  Export Results", font=("Arial", 11),
                                    bg="#1a4a6b", fg="white", relief="flat", padx=20, pady=8,
                                    cursor="hand2", state="disabled", command=self._export)
        self.export_btn.pack(side="right")

        # Handle window close gracefully during active run
        self.protocol("WM_DELETE_WINDOW", self._on_close)

    def _on_close(self):
        """Gracefully handle window close — stop worker thread before destroying."""
        if self.running:
            if not messagebox.askyesno(
                "Run In Progress",
                "A run is currently active. Progress is saved to checkpoint.\n\nAre you sure you want to quit?"
            ):
                return
            self.stop_flag = True
            self.pause_flag = False
        self.destroy()

    # ── API Key ───────────────────────────────────────────────────────────────

    def _save_api_key(self):
        key = self.api_key.get().strip()
        if not key:
            messagebox.showerror("Empty Key", "Please enter an API key before saving.")
            return
        save_api_key(key)
        self.after(0, lambda: self.key_status.config(text="✓ Key saved for next session"))

    # ── Helpers ───────────────────────────────────────────────────────────────

    def _log(self, msg, tag="dim"):
        def _do():
            self.log.configure(state="normal")
            self._log_line_count += 1
            if self._log_line_count > LOG_MAX_LINES:
                self.log.delete("1.0", "2.0")
                self._log_line_count -= 1
            self.log.insert("end", msg + "\n", tag)
            self.log.see("end")
            self.log.configure(state="disabled")
        self.after(0, _do)

    def _set_progress(self, value, label_text=None):
        def _do():
            self.progress_bar["value"] = value
            if label_text is not None:
                self.progress_label.config(text=label_text)
        self.after(0, _do)

    def _set_buttons(self, start=None, stop=None, export=None, pause=None, retry=None):
        def _do():
            if start  is not None: self.start_btn.config( state=start)
            if stop   is not None: self.stop_btn.config(  state=stop)
            if export is not None: self.export_btn.config(state=export)
            if pause  is not None: self.pause_btn.config( state=pause)
            if retry  is not None: self.retry_btn.config( state=retry)
        self.after(0, _do)

    # ── File Handling ─────────────────────────────────────────────────────────

    def _load_and_confirm(self, path, already_done=0):
        try:
            wb      = openpyxl.load_workbook(path)
            ws      = wb.active
            headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
            total   = ws.max_row - 1
        except Exception as e:
            messagebox.showerror("Error", f"Could not read Excel file:\n{e}")
            return None, None, None

        detected = detect_columns(headers)
        dialog   = ColumnConfirmDialog(self, headers, detected,
                                       total_rows=total, already_done=already_done)
        self.wait_window(dialog)

        if dialog.result is None:
            return None, None, None

        return dialog.result, dialog.chosen_delay, total

    def _new_run(self):
        path = filedialog.askopenfilename(
            title="Select Excel File",
            filetypes=[("Excel Files", "*.xlsx *.xls"), ("All Files", "*.*")]
        )
        if not path:
            return

        col_map, delay, _ = self._load_and_confirm(path, already_done=0)
        if col_map is None:
            return

        self.col_map      = col_map
        self.chosen_delay = delay

        script_dir = os.path.dirname(os.path.abspath(__file__))
        file_stem  = os.path.splitext(os.path.basename(path))[0]
        timestamp  = datetime.now().strftime("%Y-%m-%d_%H%M")
        run_name   = f"{file_stem}_{timestamp}"
        run_dir    = os.path.join(script_dir, "Runs", run_name)
        os.makedirs(run_dir, exist_ok=True)

        dest_xlsx = os.path.join(run_dir, os.path.basename(path))
        shutil.copy2(path, dest_xlsx)

        self.run_dir = run_dir
        self.run_info.config(
            text=f"Run folder: .../{os.path.join('Runs', run_name)}",
            fg="#4ddb8a"
        )
        self._log(f"New run created: {run_name}", "info")
        self._log(f"Excel copy: {os.path.basename(dest_xlsx)}", "dim")
        self.start_btn.config(state="normal")
        self.export_btn.config(state="disabled")
        self.retry_btn.config(state="disabled")

    def _resume_run(self):
        folder = filedialog.askdirectory(title="Select a Run Folder to Resume")
        if not folder:
            return

        xlsx_files = [
            f for f in os.listdir(folder)
            if f.endswith((".xlsx", ".xls")) and "Results" not in f
        ]
        if not xlsx_files:
            messagebox.showerror("Error", "No input Excel file found in that folder.")
            return

        checkpoint_path = os.path.join(folder, "checkpoint.csv")
        if not os.path.exists(checkpoint_path):
            messagebox.showwarning("Warning", "No checkpoint.csv found — this run may not have started yet.")

        done_pre  = load_checkpoint(checkpoint_path)
        xlsx_path = os.path.join(folder, xlsx_files[0])

        col_map, delay, _ = self._load_and_confirm(xlsx_path, already_done=len(done_pre))
        if col_map is None:
            return

        self.col_map      = col_map
        self.chosen_delay = delay
        self.run_dir      = folder

        self.run_info.config(
            text=f"Resuming: .../{os.path.basename(folder)}  ({len(done_pre)} rows already done)",
            fg="#f0c040"
        )
        self._log(f"Resuming run: {os.path.basename(folder)}", "warning")
        self._log(f"Already completed: {len(done_pre)} rows", "dim")
        self.start_btn.config(state="normal")
        self.export_btn.config(state="normal" if done_pre else "disabled")

        # Enable retry if there are failed rows
        summary = get_run_summary(checkpoint_path)
        self.retry_btn.config(state="normal" if summary["Errors"] > 0 else "disabled")

    # ── Run Control ───────────────────────────────────────────────────────────

    def _start(self):
        if not self.api_key.get().strip():
            messagebox.showerror("Missing API Key", "Please enter your Anthropic API key.")
            return
        if not self.run_dir:
            messagebox.showerror("No Run", "Please select New Run or Resume Run first.")
            return

        self.stop_flag  = False
        self.pause_flag = False
        self.running    = True
        self._set_buttons(start="disabled", stop="normal", export="disabled",
                          pause="normal", retry="disabled")
        threading.Thread(target=self._run_worker, daemon=True).start()

    def _stop(self):
        self.stop_flag  = True
        self.pause_flag = False
        self._log("Stop requested — finishing current business...", "warning")
        self._set_buttons(stop="disabled", pause="disabled")

    def _pause(self):
        self.pause_flag = True
        self._log("Paused — click Resume to continue.", "warning")
        self.after(0, lambda: self.pause_btn.config(
            text="▶  Resume", bg="#1a6b3a", command=self._resume_pause))

    def _resume_pause(self):
        if self.stop_flag:  # stop was hit while paused — don't resume
            return
        self.pause_flag = False
        self._log("Resumed.", "info")
        self.after(0, lambda: self.pause_btn.config(
            text="⏸  Pause", bg="#4a4a1a", command=self._pause))

    def _retry_failed(self):
        """Remove failed rows from checkpoint so they get re-checked on next run."""
        if not self.run_dir:
            return
        if self.running:
            messagebox.showwarning("Run Active", "Please stop the current run before retrying failed rows.")
            return

        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")
        done = load_checkpoint(checkpoint_path)

        error_keywords = ["Parse error", "No text response", "Error:", "Failed after"]
        good_rows = []
        retry_count = 0

        for r in done.values():
            conf     = r.get("AI_Confidence", "1")
            evidence = r.get("AI_Evidence", "")
            is_error = (conf == "0" and any(kw in evidence for kw in error_keywords))
            if is_error:
                retry_count += 1
            else:
                good_rows.append(r)

        if retry_count == 0:
            messagebox.showinfo("No Failures", "No failed rows found to retry.")
            return

        confirm = messagebox.askyesno(
            "Retry Failed Rows",
            f"Found {retry_count} failed rows.\n\n"
            f"They will be removed from the checkpoint so they get re-checked on the next run.\n\n"
            f"Continue?"
        )
        if not confirm:
            return

        rewrite_checkpoint(checkpoint_path, good_rows)
        self._log(f"Removed {retry_count} failed rows from checkpoint. Click Start to re-check them.", "warning")
        self._set_buttons(retry="disabled", start="normal")

    def _export(self):
        if not self.run_dir:
            messagebox.showerror("No Run", "No active run to export.")
            return

        xlsx_files = [
            f for f in os.listdir(self.run_dir)
            if f.endswith((".xlsx", ".xls")) and "Results" not in f
        ]
        if not xlsx_files:
            messagebox.showerror("Error", "No input Excel file found in run folder.")
            return

        input_path      = os.path.join(self.run_dir, xlsx_files[0])
        stem            = os.path.splitext(xlsx_files[0])[0]
        output_path     = os.path.join(self.run_dir, f"{stem}_Results.xlsx")
        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")

        try:
            saved = build_output_excel(input_path, output_path, checkpoint_path)
            messagebox.showinfo("Exported", f"Results saved to:\n{saved}")
            self._log(f"Exported -> {os.path.basename(saved)}", "success")
        except PermissionError:
            messagebox.showerror(
                "File In Use",
                "Could not save — the Results file may be open in Excel.\n"
                "Please close it and try again."
            )
        except Exception as e:
            messagebox.showerror("Export Error", str(e))

    # ── Background Worker ─────────────────────────────────────────────────────

    def _run_worker(self):
        api_key = self.api_key.get().strip()

        xlsx_files = [
            f for f in os.listdir(self.run_dir)
            if f.endswith((".xlsx", ".xls")) and "Results" not in f
        ]
        if not xlsx_files:
            self._log("No Excel file found in run folder.", "error")
            self._set_buttons(start="normal", stop="disabled", pause="disabled")
            return

        input_path      = os.path.join(self.run_dir, xlsx_files[0])
        checkpoint_path = os.path.join(self.run_dir, "checkpoint.csv")

        try:
            client     = anthropic.Anthropic(api_key=api_key)
            wb         = openpyxl.load_workbook(input_path)
            ws         = wb.active
            total_rows = ws.max_row - 1
            done       = load_checkpoint(checkpoint_path)
            col        = self.col_map

            self._log(f"File: {os.path.basename(input_path)}", "info")
            self._log(f"Total: {total_rows}  |  Done: {len(done)}  |  Remaining: {total_rows - len(done)}", "info")
            self._log("─" * 55, "dim")

            self.after(0, lambda: self.progress_bar.config(maximum=total_rows, value=len(done)))

            for row_idx in range(2, ws.max_row + 1):
                if self.stop_flag:
                    self._log("Stopped. Progress saved to checkpoint.", "warning")
                    break

                if self.pause_flag:
                    self._log(f"Paused at row {row_idx - 1}. Click Resume to continue.", "warning")
                    while self.pause_flag and not self.stop_flag:
                        time.sleep(0.5)
                    if self.stop_flag:
                        self._log("Stopped during pause.", "warning")
                        break
                    self._log(f"Resumed from row {row_idx - 1}.", "info")

                key = str(row_idx)
                if key in done:
                    self._set_progress(row_idx - 1,
                                       f"Skipping already-done rows... ({row_idx - 1}/{total_rows})")
                    continue

                name    = get_cell_value(ws, row_idx, col.get("name"))
                website = get_cell_value(ws, row_idx, col.get("website"))
                city    = get_cell_value(ws, row_idx, col.get("city"))
                state   = get_cell_value(ws, row_idx, col.get("state"))

                progress = row_idx - 1
                self._set_progress(progress, f"Checking {progress} of {total_rows}: {name}")
                self._log(f"[{progress}/{total_rows}] {name} ({city}, {state})", "dim")

                status, confidence, evidence = check_business(client, name, website, city, state)

                checked_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                tag = "success" if status == "Active" else \
                      "error"   if status == "Likely Closed" else "warning"
                self._log(f"  -> {status} ({confidence}%) | {evidence}", tag)

                save_checkpoint(checkpoint_path, key, name, website,
                                status, confidence, evidence, checked_at)

                time.sleep(self.chosen_delay)

            if not self.stop_flag:
                self._log("─" * 55, "dim")
                self._log("All businesses checked!", "success")
                self._set_progress(total_rows, "Complete!")

                # Show summary dialog
                summary = get_run_summary(checkpoint_path)
                self.after(0, lambda: SummaryDialog(
                    self, summary,
                    on_export=self._export,
                    on_retry=self._retry_failed
                ))

                # Enable retry button if there were errors
                if summary["Errors"] > 0:
                    self._set_buttons(retry="normal")

        except Exception as e:
            self._log(f"Fatal error: {str(e)}", "error")
            err = str(e)  # capture before 'e' is deleted at end of except block
            self.after(0, lambda: messagebox.showerror("Error", err))

        finally:
            self.running = False
            self._set_buttons(start="normal", stop="disabled", pause="disabled", export="normal")


if __name__ == "__main__":
    app = App()
    app.mainloop()
