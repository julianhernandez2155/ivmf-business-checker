"""
Business Operational Status Checker
Uses Claude AI with web search to determine if businesses are still operating.

SETUP:
1. pip install anthropic openpyxl
2. In PowerShell: $env:ANTHROPIC_API_KEY="sk-ant-your-key-here"
3. Place BMOSG_All_Businesses.xlsx in the same folder
4. Delete checkpoint.csv if starting fresh
5. Run: python business_checker.py
"""

import os
import csv
import time
import json
import openpyxl
from datetime import datetime
import anthropic

# ── Configuration ────────────────────────────────────────────────────────────

INPUT_FILE     = "BMOSG_All_Businesses.xlsx"
OUTPUT_FILE    = "BMOSG_Results.xlsx"
CHECKPOINT_CSV = "checkpoint.csv"
DELAY_SECONDS  = 8
MAX_RETRIES    = 4
MODEL          = "claude-haiku-4-5-20251001"

COL_NAME    = 0
COL_WEBSITE = 5
COL_CITY    = 9
COL_STATE   = 10

# ─────────────────────────────────────────────────────────────────────────────

def load_checkpoint(path):
    done = {}
    if not os.path.exists(path):
        return done
    with open(path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for row in reader:
            done[row["row_index"]] = row
    print(f"[Checkpoint] Resuming - {len(done)} rows already completed.")
    return done


def save_checkpoint(path, row_index, name, website, status, confidence, evidence, checked_at):
    file_exists = os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=[
            "row_index", "name", "website",
            "AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"
        ])
        if not file_exists:
            writer.writeheader()
        writer.writerow({
            "row_index":      row_index,
            "name":           name,
            "website":        website,
            "AI_Status":      status,
            "AI_Confidence":  confidence,
            "AI_Evidence":    evidence,
            "AI_Checked_At":  checked_at,
        })


def check_business(client, name, website, city, state):
    location     = ", ".join(filter(None, [city, state]))
    website_note = f"Their listed website is: {website}" if website else "No website listed."

    prompt = f"""Search the web for "{name}" located in {location} and determine if this business is still actively operating.

Check their website ({website_note}), Google, Yelp, or any other sources for signs of active operation or closure.

After searching, you MUST respond with ONLY this JSON and nothing else — no intro, no explanation, no markdown:
{{"status": "Active", "confidence": 85, "evidence": "Website is live with recent activity."}}

Status must be one of: Active, Likely Closed, Uncertain, No Web Presence
Confidence must be an integer 0-100."""

    for attempt in range(MAX_RETRIES):
        try:
            messages = [{"role": "user", "content": prompt}]

            response = client.messages.create(
                model=MODEL,
                max_tokens=200,
                tools=[{"type": "web_search_20250305", "name": "web_search"}],
                messages=messages
            )

            # If Claude stopped mid-search without a text summary, nudge it to finalize
            if response.stop_reason == "tool_use":
                messages.append({"role": "assistant", "content": response.content})
                messages.append({"role": "user", "content": "Now respond with ONLY the JSON result. No explanation, no markdown, just the JSON object."})
                response = client.messages.create(
                    model=MODEL,
                    max_tokens=200,
                    tools=[{"type": "web_search_20250305", "name": "web_search"}],
                    messages=messages
                )

            # Collect ALL text blocks — we want the LAST one (after search completes)
            text_blocks = []
            for block in response.content:
                if hasattr(block, "type") and block.type == "text":
                    text_blocks.append(block.text.strip())

            if not text_blocks:
                return ("Uncertain", "0", "No text response from model.")

            # Use the last text block — final answer after tool use
            result_text = text_blocks[-1]

            # Strip markdown fences if present
            result_text = result_text.replace("```json", "").replace("```", "").strip()

            # Extract JSON object even if there is surrounding text
            start = result_text.find("{")
            end   = result_text.rfind("}") + 1
            if start != -1 and end > start:
                result_text = result_text[start:end]

            parsed = json.loads(result_text)
            return (
                parsed.get("status", "Uncertain"),
                str(parsed.get("confidence", 50)),
                parsed.get("evidence", "No evidence found.")
            )

        except json.JSONDecodeError:
            return ("Uncertain", "0", f"Parse error: {result_text[:100]}")

        except anthropic.RateLimitError:
            wait = 30 * (attempt + 1)
            print(f"         [Rate limited] Waiting {wait}s (attempt {attempt + 1}/{MAX_RETRIES})...")
            time.sleep(wait)

        except Exception as e:
            return ("Uncertain", "0", f"Error: {str(e)[:100]}")

    return ("Uncertain", "0", "Failed after max retries.")


def build_output_excel(input_path, output_path, checkpoint_path):
    print(f"\n[Output] Building {output_path} ...")
    checkpoint = load_checkpoint(checkpoint_path)

    wb_in  = openpyxl.load_workbook(input_path)
    ws_in  = wb_in.active
    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Results"

    headers = [ws_in.cell(1, c).value for c in range(1, ws_in.max_column + 1)]
    headers += ["AI_Status", "AI_Confidence", "AI_Evidence", "AI_Checked_At"]
    ws_out.append(headers)

    for row_idx in range(2, ws_in.max_row + 1):
        row_data = [ws_in.cell(row_idx, c).value for c in range(1, ws_in.max_column + 1)]
        key = str(row_idx)
        if key in checkpoint:
            r = checkpoint[key]
            row_data += [r["AI_Status"], r["AI_Confidence"], r["AI_Evidence"], r["AI_Checked_At"]]
        else:
            row_data += ["Not Checked", "", "", ""]
        ws_out.append(row_data)

    wb_out.save(output_path)
    print(f"[Output] Saved -> {output_path}")


def main():
    api_key = os.environ.get("ANTHROPIC_API_KEY")
    if not api_key:
        print("ERROR: ANTHROPIC_API_KEY not set.")
        print('Run this in PowerShell first: $env:ANTHROPIC_API_KEY="sk-ant-your-key-here"')
        return

    client = anthropic.Anthropic(api_key=api_key)

    wb = openpyxl.load_workbook(INPUT_FILE)
    ws = wb.active
    total_rows = ws.max_row - 1

    done = load_checkpoint(CHECKPOINT_CSV)

    print(f"\n{'='*55}")
    print(f"  Business Operational Status Checker")
    print(f"  Total businesses : {total_rows}")
    print(f"  Already checked  : {len(done)}")
    print(f"  Remaining        : {total_rows - len(done)}")
    print(f"  Delay per call   : {DELAY_SECONDS}s")
    print(f"{'='*55}\n")

    for row_idx in range(2, ws.max_row + 1):
        key = str(row_idx)

        if key in done:
            continue

        name    = ws.cell(row_idx, COL_NAME + 1).value or ""
        website = ws.cell(row_idx, COL_WEBSITE + 1).value or ""
        city    = ws.cell(row_idx, COL_CITY + 1).value or ""
        state   = ws.cell(row_idx, COL_STATE + 1).value or ""

        progress = row_idx - 1
        print(f"[{progress}/{total_rows}] Checking: {name} ({city}, {state})")

        status, confidence, evidence = check_business(client, name, website, city, state)
        checked_at = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        print(f"         -> {status} ({confidence}%) | {evidence}")

        save_checkpoint(CHECKPOINT_CSV, key, name, website, status, confidence, evidence, checked_at)

        time.sleep(DELAY_SECONDS)

    print(f"\n[Done] All rows processed.")
    build_output_excel(INPUT_FILE, OUTPUT_FILE, CHECKPOINT_CSV)
    print(f"\nResults saved to: {OUTPUT_FILE}")
    print(f"Checkpoint saved to: {CHECKPOINT_CSV}\n")


if __name__ == "__main__":
    main()
